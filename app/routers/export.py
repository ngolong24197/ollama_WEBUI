import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Conversation, Message
from app.schemas.conversation import ConversationResponse, MessageResponse

router = APIRouter(prefix="/api/conversations", tags=["export"])

ROLE_LABELS = {"user": "You:", "assistant": "Assistant:"}


def _add_code_block(doc: Document, code: str):
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Courier New"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is not None:
        rFonts.set(qn("w:ascii"), "Courier New")
        rFonts.set(qn("w:hAnsi"), "Courier New")
    p.paragraph_format.space_before = 4
    p.paragraph_format.space_after = 4


def _build_document(conv: Conversation, messages: list[Message]) -> bytes:
    doc = Document()

    # Header with model name and date
    header = doc.sections[0].header
    header_p = header.paragraphs[0]
    header_p.text = f"Model: {conv.model}  |  Exported: {datetime.now():%Y-%m-%d %H:%M}"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Title
    title = doc.add_heading(conv.title, level=0)

    # Messages
    for msg in messages:
        label = ROLE_LABELS.get(msg.role, f"{msg.role.capitalize()}:")
        p = doc.add_paragraph()
        role_run = p.add_run(label + " ")
        role_run.bold = True
        content = msg.content or ""

        # Split content by code fence markers
        parts = content.split("```")
        for i, part in enumerate(parts):
            if not part:
                continue
            if i % 2 == 0:
                # Regular text
                if part.strip():
                    p.add_run(part)
            else:
                # Inside a code block – strip optional language tag on first line
                lines = part.split("\n", 1)
                code = lines[1] if len(lines) > 1 and lines[0].strip().isidentifier() else part
                p = _add_code_block(doc, code.strip())

    # Token count summary
    doc.add_paragraph()
    summary = doc.add_paragraph()
    total_prompt = sum(m.prompt_tokens for m in messages)
    total_completion = sum(m.completion_tokens for m in messages)
    summary_run = summary.add_run(
        f"Token Summary — Prompt: {total_prompt:,} | "
        f"Completion: {total_completion:,} | "
        f"Total: {total_prompt + total_completion:,}"
    )
    summary_run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


@router.get("/{conversation_id}/export")
async def export_conversation(
    conversation_id: int,
    format: str = "docx",
    db: AsyncSession = Depends(get_db),
):
    conv = await db.get(Conversation, conversation_id)
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found")

    result = await db.execute(
        select(Message)
        .where(Message.conversation_id == conversation_id)
        .order_by(Message.created_at)
    )
    messages = list(result.scalars().all())

    try:
        content = _build_document(conv, messages)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Export failed: {exc}") from exc

    filename = f"{conv.title[:50].replace(' ', '_')}.docx"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )